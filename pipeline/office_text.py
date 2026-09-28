"""Текст документов Word и Excel прежних форматов: DOC, XLS, XLSX. Задача #112.

DOCX читается без зависимостей (`matrix_rules.docx_text`): это zip с XML. У остальных форматов
своего разбора нет, и до #112 они лежали в реестре со статусом «формат не читается»: у Полярной 25,
ДОО — 162 DOC, 22 XLS и 10 XLSX, в основном в исполнительной документации (акты приёмки систем,
акты испытаний, паспорта вентсистем, реестры передаваемой документации).

- **DOC** — внешней программой. В образе воркера это `catdoc` (Debian trixie, пакет catdoc):
  он отдаёт текст в порядке чтения, ячейки таблицы — через табуляцию. `antiword` рисует таблицы
  псевдографикой и рвёт фразу по ячейкам («ответственного | | за эксплуатацию»), поэтому он только
  запасной. На macOS есть `textutil`. Какая программа нашлась — видно в `converter()`.
- **XLS** — `xlrd` (он читает только старый формат), **XLSX** — `openpyxl`. Строка листа — строка
  текста, ячейки через табуляцию, листы — через пустую строку с названием листа. Целые числа пишутся
  без «.0», дроби — с запятой, как в документах; даты — «дд.мм.гггг».

Весь текст документа отдаётся одной страницей, как у DOCX (#42): у этих форматов страниц нет.
"""
import datetime
import os
import shutil
import subprocess

# Программы для DOC в порядке предпочтения: (команда, аргументы до пути)
DOC_CONVERTERS = (
    ("catdoc", ["-w", "-d", "utf-8"]),        # -w: без переноса строк, иначе фразы рвутся на 72-м знаке
    ("textutil", ["-convert", "txt", "-stdout"]),
    ("antiword", ["-m", "UTF-8.txt"]),
)
TIMEOUT_S = 60
OFFICE_EXT = (".doc", ".xls", ".xlsx")


def converter():
    """Программа, которой будет прочитан DOC, или None: её нет ни одной."""
    return next((name for name, _ in DOC_CONVERTERS if shutil.which(name)), None)


def doc_text(path):
    """Текст DOC. Пустая строка — файл не прочитался или программы для DOC нет."""
    for name, args in DOC_CONVERTERS:
        if not shutil.which(name):
            continue
        try:
            done = subprocess.run([name, *args, path], capture_output=True, timeout=TIMEOUT_S)
        except (OSError, subprocess.TimeoutExpired):
            continue
        text = done.stdout.decode("utf-8", "replace")
        if text.strip():
            return text
    return ""


def cell_text(value):
    """Значение ячейки как в документе: 12.0 → «12», 3.5 → «3,5», дата → «дд.мм.гггг»."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime.datetime):
        return value.strftime("%d.%m.%Y") if (value.hour, value.minute, value.second) == (0, 0, 0) \
            else value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(round(value, 6)).replace(".", ",")
    return str(value).strip()


def _rows_text(title, rows):
    lines = [f"Лист: {title}"]
    for row in rows:
        cells = [cell_text(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append("\t".join(cells))
    return "\n".join(lines)


def xls_text(path):
    """Текст XLS по листам."""
    try:
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
    except Exception:
        return ""
    parts = []
    try:
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except Exception:
                            pass
                    row.append(cell.value if cell.value != "" else None)
                rows.append(row)
            parts.append(_rows_text(sheet.name, rows))
    finally:
        book.release_resources()
    return "\n\n".join(parts)


def xlsx_text(path):
    """Текст XLSX по листам: значения формул — сохранённые в файле, не пересчитанные."""
    try:
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return ""
    try:
        return "\n\n".join(_rows_text(ws.title, ws.iter_rows(values_only=True)) for ws in book.worksheets)
    finally:
        book.close()


def text(path):
    """Текст документа по расширению: DOC, XLS, XLSX. Для прочих — пустая строка."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".doc":
        return doc_text(path)
    if ext == ".xls":
        return xls_text(path)
    if ext == ".xlsx":
        return xlsx_text(path)
    return ""
